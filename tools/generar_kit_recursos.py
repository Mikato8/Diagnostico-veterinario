#!/usr/bin/env python3
"""Genera las plantillas DOCX del kit de recursos para veterinarios."""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parents[1]
OUT = BASE / "public" / "descargas"
OUT.mkdir(parents=True, exist_ok=True)

TEAL = RGBColor(0x0D, 0x5C, 0x4D)
INK = RGBColor(0x12, 0x21, 0x1D)
MUTED = RGBColor(0x64, 0x74, 0x6E)
SAND = "F7F3EA"
TEAL_LIGHT = "DFF2EA"


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: RGBColor | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.space_before = Pt(2)


def setup(document: Document, title: str, subtitle: str) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size in [("Title", 24), ("Heading 1", 15), ("Heading 2", 12)]:
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.color.rgb = TEAL
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.text = "DIAGNÓSTICO VETERINARIO · KIT PARA VETERINARIOS"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED
        run.font.bold = True

    footer = section.footer.paragraphs[0]
    footer.text = "Documento educativo · No sustituye el juicio clínico ni la normativa aplicable."
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED

    document.add_paragraph(title, style="Title")
    paragraph = document.add_paragraph(subtitle)
    paragraph.paragraph_format.space_after = Pt(14)
    for run in paragraph.runs:
        run.font.size = Pt(11)
        run.font.color.rgb = MUTED
        run.italic = True


def label_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], label, bold=True)
        shade(cells[0], SAND)
        set_cell_text(cells[1], value)
    document.add_paragraph()


def field_block(document: Document, title: str, hint: str = "", lines: int = 4) -> None:
    document.add_heading(title, level=1)
    if hint:
        paragraph = document.add_paragraph(hint)
        paragraph.paragraph_format.space_after = Pt(4)
        for run in paragraph.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = MUTED
            run.italic = True
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    cell.text = "\n" * lines
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.2
    document.add_paragraph()


def bullet_list(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(item, style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)


def checkbox_table(document: Document, items: list[str]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for item in items:
        cells = table.add_row().cells
        set_cell_text(cells[0], "☐", bold=True)
        cells[0].width = Cm(1.0)
        set_cell_text(cells[1], item)
    document.add_paragraph()


def build_case_template() -> Path:
    document = Document()
    setup(
        document,
        "Plantilla para documentar un caso clínico",
        "Úsala como guía de redacción. Antes de publicar, elimina nombres, direcciones y cualquier dato que permita identificar al tutor o al paciente.",
    )

    document.add_heading("1. Identificación del caso", level=1)
    label_table(document, [
        ("Título propuesto", "Describe especie, problema y contexto en una frase."),
        ("Especie", ""),
        ("Raza", ""),
        ("Sexo", "Macho / Hembra / No determinado"),
        ("Edad", ""),
        ("Peso", ""),
        ("Estado de México", ""),
        ("Entorno", "Clínica / Hospital / Campo / Laboratorio / Otro"),
        ("Autor", ""),
        ("Institución", ""),
    ])

    field_block(document, "2. Motivo de consulta", "¿Qué motivó la consulta? Incluye duración y evolución.", 2)
    field_block(document, "3. Anamnesis", "Antecedentes, manejo, alimentación, vacunación, exposición, tratamientos previos. Sin datos identificables.", 6)
    field_block(document, "4. Exploración física", "Constantes, condición corporal y hallazgos por aparatos y sistemas.", 6)
    field_block(document, "5. Estudios y resultados", "Hemograma, bioquímica, imagen, citología, cultivo, PCR u otros. Incluye unidades y valores de referencia.", 6)
    field_block(document, "6. Diagnósticos diferenciales", "Enumera las hipótesis principales y cómo las descartaste.", 3)
    field_block(document, "7. Diagnóstico final", "Escribe el diagnóstico sustentado. Si aplica, agrega código ICD-11.", 2)
    field_block(document, "8. Tratamiento y manejo", "Describe el manejo, la respuesta y los ajustes. Evita dosis que puedan interpretarse como prescripción universal.", 6)
    field_block(document, "9. Desenlace", "Resuelto / En tratamiento / Con secuela / Fallecido / Eutanasia / Sin seguimiento.", 2)
    field_block(document, "10. Discusión y aprendizaje", "¿Qué harías diferente? ¿Qué limitaciones tuvo el caso? ¿Qué mensaje deja para la práctica?", 6)
    field_block(document, "11. Referencias", "Cita libros, guías o artículos consultados.", 3)

    document.add_heading("Antes de enviar", level=1)
    document.add_paragraph("Marca cada punto. Si falta uno, el caso no debe publicarse.")
    checkbox_table(document, [
        "Cuento con consentimiento informado del tutor para publicar con fines educativos.",
        "Eliminé nombres, teléfonos, correos, direcciones, números de expediente y microchips.",
        "Revisé que las imágenes no muestren personas, documentos ni datos identificables.",
        "El diagnóstico final está sustentado por la historia, la exploración y los estudios.",
        "No incluí dosis o indicaciones que puedan interpretarse como prescripción.",
        "Declaré conflictos de interés si el caso promueve un producto o servicio.",
        "Elegí una licencia de publicación y entiendo que el caso podrá ser citado.",
    ])

    path = OUT / "plantilla-caso-clinico.docx"
    document.save(path)
    return path


def build_consent_checklist() -> Path:
    document = Document()
    setup(
        document,
        "Checklist de consentimiento y anonimización",
        "Imprime esta lista y revísala antes de publicar un caso en el portal o compartirlo con colegas.",
    )
    document.add_paragraph("Este documento no sustituye el consentimiento informado de tu institución. Úsalo como verificación interna y conserva el consentimiento firmado en el expediente.")
    document.add_heading("Consentimiento del tutor", level=1)
    checkbox_table(document, [
        "Expliqué al tutor que el caso será público y podrá ser citado.",
        "Expliqué que no se publicarán datos que identifiquen a la persona o al animal.",
        "El tutor autorizó el uso educativo del caso y de las imágenes.",
        "El tutor sabe que puede solicitar la corrección o retiro del contenido.",
        "Conservo el consentimiento firmado o el registro institucional correspondiente.",
    ])

    document.add_heading("Anonimización del texto", level=1)
    checkbox_table(document, [
        "Eliminé el nombre del tutor y del paciente.",
        "Eliminé direcciones, teléfonos, correos y redes sociales.",
        "Eliminé números de expediente, folios internos y datos de facturación.",
        "Eliminé nombres de personas del equipo cuando no sean relevantes.",
        "Revisé que la combinación de edad, raza, colonia y fecha no permita identificar al caso.",
    ])

    document.add_heading("Imágenes", level=1)
    checkbox_table(document, [
        "No aparecen personas, gafetes, documentos ni pantallas con datos.",
        "No se ven placas, microchips, collares con nombre ni señas particulares.",
        "Recorté o difuminé cualquier identificador visual.",
        "Las imágenes tienen relación clínica y no muestran sufrimiento innecesario.",
        "Cuento con autorización para publicar las imágenes.",
    ])

    document.add_heading("Publicación", level=1)
    checkbox_table(document, [
        "El caso no contiene publicidad encubierta ni promoción de productos.",
        "El diagnóstico está sustentado y las limitaciones se explican con honestidad.",
        "Elegí la licencia de publicación (por defecto CC BY-NC-SA 4.0).",
        "Revisé el texto antes de enviarlo a revisión editorial.",
    ])

    document.add_heading("Firmas de verificación interna", level=1)
    label_table(document, [
        ("Revisó (nombre)", ""),
        ("Cargo", ""),
        ("Fecha", ""),
        ("Firma", ""),
    ])

    path = OUT / "checklist-consentimiento.docx"
    document.save(path)
    return path


def build_request_template() -> Path:
    document = Document()
    setup(
        document,
        "Solicitud de acceso a datos veterinarios",
        "Carta modelo para pedir acceso a un dataset, repositorio o sistema de vigilancia. Completa los campos entre corchetes.",
    )
    document.add_paragraph("[Ciudad], [fecha]")
    document.add_paragraph("A: [nombre del equipo o institución responsable]")
    document.add_paragraph("Asunto: Solicitud de acceso a [nombre del recurso]")

    document.add_heading("Presentación", level=1)
    document.add_paragraph("Mi nombre es [nombre completo]. Soy médico veterinario con cédula profesional [número] y me desempeño en [institución o práctica]. Mi correo es [correo] y mi teléfono es [teléfono].")

    document.add_heading("Objetivo", level=1)
    document.add_paragraph("Solicito acceso a [nombre del recurso] para [describir el objetivo: tesis, artículo, docencia, análisis de casos, vigilancia]. La pregunta principal es [pregunta concreta].")

    document.add_heading("Uso previsto", level=1)
    document.add_paragraph("Planeo utilizar los datos para [análisis descriptivo, comparación, validación, docencia]. No los usaré para identificar personas, animales o establecimientos, ni para decisiones individuales de diagnóstico.")

    document.add_heading("Compromisos", level=1)
    bullet_list(document, [
        "Cumpliré las condiciones de acceso y la licencia del recurso.",
        "No compartiré los datos con terceros sin autorización.",
        "Conservaré los archivos en un entorno seguro y los eliminaré cuando termine el proyecto.",
        "Citaré la fuente y el DOI o referencia indicada por ustedes.",
        "Enviaré una copia del producto final o un resumen de resultados si lo solicitan.",
    ])

    document.add_heading("Datos que necesito", level=1)
    label_table(document, [
        ("Variables o archivos", ""),
        ("Periodo", ""),
        ("Especie o población", ""),
        ("Formato deseado", "CSV / Excel / PDF / Otro"),
        ("¿Requiere diccionario de variables?", "Sí / No"),
    ])

    document.add_heading("Cierre", level=1)
    document.add_paragraph("Quedo atento a los requisitos, tiempos y documentos que deba firmar. Agradezco su orientación para usar el recurso de forma ética y correcta.")
    document.add_paragraph("Saludos cordiales,")
    document.add_paragraph("[nombre completo]")
    document.add_paragraph("[institución]")
    document.add_paragraph("[correo] · [teléfono]")

    document.add_heading("Anexo: verificación ética", level=1)
    checkbox_table(document, [
        "El proyecto cuenta con aprobación de comité de ética o investigación, si aplica.",
        "Tengo un plan para proteger la confidencialidad de los datos.",
        "Entiendo que el acceso puede ser negado o condicionado.",
        "No usaré los datos para reidentificar personas o establecimientos.",
    ])

    path = OUT / "solicitud-acceso-datos.docx"
    document.save(path)
    return path


def build_guide() -> Path:
    document = Document()
    setup(
        document,
        "Guía rápida: recursos para veterinarios",
        "Qué es un dataset, un repositorio y una API, y cómo pedir acceso sin ser ingeniero en software.",
    )
    document.add_heading("No necesitas programar para usar estos recursos", level=1)
    document.add_paragraph("Muchos recursos veterinarios viven en repositorios de código o de datos. Eso no significa que tengas que instalar nada. En la mayoría de los casos basta con leer la ficha, descargar un archivo o enviar una solicitud. El portal te explica cada paso en español.")

    document.add_heading("Tres conceptos básicos", level=1)
    label_table(document, [
        ("Dataset", "Un conjunto de datos: filas y columnas, imágenes o registros. Se abre con Excel, Numbers o Google Sheets."),
        ("Repositorio", "Un sitio donde se guardan y comparten archivos. GitHub aloja código; Dryad aloja datos de investigación."),
        ("API", "Una forma de conectar sistemas. No necesitas programar para leer sobre ella, pero sí para integrarla en tu software."),
    ])

    document.add_heading("Cómo elegir un recurso", level=1)
    bullet_list(document, [
        "Define tu objetivo: docencia, tesis, vigilancia, comparación de casos o mejora de tu sistema.",
        "Revisa la especie, el país y el periodo de los datos.",
        "Lee la licencia y las condiciones de acceso antes de descargar.",
        "Verifica quién publicó el recurso y si tiene una revisión editorial.",
        "Si no entiendes un término, consulta el glosario del portal.",
    ])

    document.add_heading("Cómo pedir acceso en cinco pasos", level=1)
    for index, item in enumerate([
        "Abre la ficha interna del recurso y lee los requisitos.",
        "Copia la plantilla de solicitud y completa tus datos.",
        "Explica tu objetivo con una pregunta concreta.",
        "Menciona cómo protegerás la confidencialidad y qué harás con los resultados.",
        "Guarda la respuesta y las condiciones de uso en tu expediente del proyecto.",
    ], start=1):
        document.add_paragraph(f"{index}. {item}")

    document.add_heading("Señales de alerta", level=1)
    bullet_list(document, [
        "El recurso no indica licencia, autor o fecha.",
        "Piden datos personales de pacientes o tutores sin justificación.",
        "Prometen acceso inmediato a cambio de pagos informales.",
        "No permiten citar la fuente ni explican cómo se generaron los datos.",
    ])

    document.add_heading("Antes de usar datos en una investigación", level=1)
    checkbox_table(document, [
        "Tengo una pregunta de investigación clara.",
        "Revisé la licencia y las condiciones de acceso.",
        "Contemplé la aprobación de un comité de ética si corresponde.",
        "Tengo un plan para resguardar los archivos.",
        "Sé cómo voy a citar la fuente.",
        "No voy a reidentificar personas, animales o establecimientos.",
    ])

    document.add_heading("Dónde pedir ayuda", level=1)
    document.add_paragraph("Consulta la sección de recursos del portal, el glosario y las plantillas descargables. Si trabajas con una institución, pregunta por la oficina de investigación o biblioteca; suelen tener apoyo para búsqueda de datos y gestión de información.")

    path = OUT / "guia-recursos-para-veterinarios.docx"
    document.save(path)
    return path


def main() -> None:
    files = [
        build_case_template(),
        build_consent_checklist(),
        build_request_template(),
        build_guide(),
    ]
    zip_path = OUT / "kit-recursos-veterinarios.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for file in files:
            archive.write(file, arcname=file.name)
    for file in files + [zip_path]:
        print(f"{file.relative_to(BASE)} · {file.stat().st_size} bytes")


if __name__ == "__main__":
    main()
